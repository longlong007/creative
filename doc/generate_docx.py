#!/usr/bin/env python3
"""Convert doc/*.md into .docx files in the same directory."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1] / "doc"


def set_run_font(run, *, east_asia: str = "微软雅黑", ascii_font: str = "Calibri", size: int | None = None, bold=None):
    run.font.name = ascii_font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east_asia)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_formatted_paragraph(doc, text: str, style: str | None = None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pattern = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = p.add_run(text[pos : match.start()])
            set_run_font(run, size=11)
        token = match.group(0)
        if token.startswith("`"):
            run = p.add_run(token.strip("`"))
            set_run_font(run, ascii_font="Consolas", east_asia="微软雅黑", size=10)
        else:
            run = p.add_run(token.strip("*"))
            set_run_font(run, size=11, bold=True)
        pos = match.end()
    if pos < len(text):
        run = p.add_run(text[pos:])
        set_run_font(run, size=11)
    if not text:
        set_run_font(p.add_run(""), size=11)
    return p


def parse_table(lines: list[str]) -> list[list[str]] | None:
    if len(lines) < 2:
        return None
    if not all("|" in line for line in lines):
        return None
    rows = []
    for i, line in enumerate(lines):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if i == 1 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows or None


def add_heading_styled(doc, text: str, level: int):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 18, 2: 16, 3: 14, 4: 12}.get(level, 12), bold=True)
    return h


def convert_markdown(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    i = 0
    in_code = False
    code_buf: list[str] = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                set_run_font(run, ascii_font="Consolas", east_asia="微软雅黑", size=9)
                p.paragraph_format.space_after = Pt(8)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith("# "):
            add_heading_styled(doc, line[2:].strip(), 1)
            i += 1
            continue
        if line.startswith("## "):
            add_heading_styled(doc, line[3:].strip(), 2)
            i += 1
            continue
        if line.startswith("### "):
            add_heading_styled(doc, line[4:].strip(), 3)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append(lines[i])
                i += 1
            rows = parse_table(block)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        cell_obj = table.cell(r_idx, c_idx)
                        cell_obj.text = ""
                        run = cell_obj.paragraphs[0].add_run(cell)
                        set_run_font(run, size=10, bold=(r_idx == 0))
            continue
        if re.match(r"^\d+\.\s", line.strip()):
            p = add_formatted_paragraph(doc, re.sub(r"^\d+\.\s", "", line.strip()), style="List Number")
            i += 1
            continue
        if line.strip().startswith("- "):
            add_formatted_paragraph(doc, line.strip()[2:], style="List Bullet")
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        add_formatted_paragraph(doc, line.strip())
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)


def main() -> None:
    files = sorted(ROOT.glob("*.md"))
    if not files:
        raise SystemExit(f"no markdown in {ROOT}")
    for md in files:
        out = md.with_suffix(".docx")
        convert_markdown(md, out)
        print(f"wrote {out.name}")


if __name__ == "__main__":
    main()
